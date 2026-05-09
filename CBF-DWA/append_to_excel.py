import pandas as pd
import openpyxl
import numpy as np


def app_to_xlsx(data, filename="sac.xlsx"):
    # 找到数据中最长数组的长度
    max_len = max(len(v) if isinstance(v, (list, np.ndarray)) else 1 for v in data.values())

    # 填充所有较短的数组，确保它们的长度与最大长度一致
    for key in data:
        if isinstance(data[key], (list, np.ndarray)) and len(data[key]) < max_len:
            # 填充NaN，直到长度与最大长度一致
            data[key] = np.pad(data[key], (0, max_len - len(data[key])), constant_values=np.nan)

    # 将处理后的数据转换为DataFrame
    df = pd.DataFrame(data)

    # 打开Excel文件，如果文件不存在则创建
    try:
        wb = openpyxl.load_workbook(filename)
    except FileNotFoundError:
        wb = openpyxl.Workbook()

    sheet = wb.active

    # 获取当前工作表的最大列数
    max_col = sheet.max_column

    # 如果工作表为空，设置最大列数为0
    if max_col == 1 and sheet.cell(row=1, column=1).value is None:
        max_col = 0

    # 新数据的起始列位置
    start_col = max_col + 1

    # 将 DataFrame 按列写入到 Excel
    for col_num, column in enumerate(df.columns, start=start_col):
        for row_num, value in enumerate(df[column], start=1):
            sheet.cell(row=row_num, column=col_num, value=value)

    # 保存 Excel 文件
    wb.save(filename)
